# encoding: utf-8

import os

from behave import given, when, then
from hamcrest import assert_that, is_, is_not, greater_than
from StringIO import StringIO

from pptx import Presentation


def absjoin(*paths):
    return os.path.abspath(os.path.join(*paths))

thisdir = os.path.split(__file__)[0]
scratch_dir = absjoin(thisdir, '../_scratch')
test_file_dir = absjoin(thisdir, '../../tests/test_files')
basic_pptx_path = absjoin(test_file_dir, 'test.pptx')
saved_pptx_path = absjoin(scratch_dir, 'test_out.pptx')


# given ===================================================

@given('a clean working directory')
def step_given_clean_working_dir(context):
    if os.path.isfile(saved_pptx_path):
        os.remove(saved_pptx_path)


@given('an initialized pptx environment')
def step_given_initialized_pptx_env(context):
    pass


@given('I have an empty presentation open')
def step_given_empty_prs(context):
    context.prs = Presentation()


# when ====================================================

@when('I construct a Presentation instance with no path argument')
def step_when_construct_default_prs(context):
    context.prs = Presentation()


@when('I open a basic PowerPoint presentation')
def step_when_open_basic_pptx(context):
    context.prs = Presentation(basic_pptx_path)


@when('I open a presentation contained in a stream')
def step_when_open_presentation_stream(context):
    with open(basic_pptx_path) as f:
        stream = StringIO(f.read())
    context.prs = Presentation(stream)
    stream.close()


@when('I save that stream to a file')
def step_when_save_stream_to_a_file(context):
    if os.path.isfile(saved_pptx_path):
        os.remove(saved_pptx_path)
    context.stream.seek(0)
    with open(saved_pptx_path, 'wb') as f:
        f.write(context.stream.read())


@when('I save the presentation')
def step_when_save_presentation(context):
    if os.path.isfile(saved_pptx_path):
        os.remove(saved_pptx_path)
    context.prs.save(saved_pptx_path)


@when('I save the presentation to a stream')
def step_when_save_presentation_to_stream(context):
    context.stream = StringIO()
    context.prs.save(context.stream)


# then ====================================================

@then('I receive a presentation based on the default template')
def step_then_receive_prs_based_on_def_tmpl(context):
    prs = context.prs
    assert_that(prs, is_not(None))
    slidemasters = prs.slidemasters
    assert_that(slidemasters, is_not(None))
    assert_that(len(slidemasters), is_(1))
    slidelayouts = slidemasters[0].slidelayouts
    assert_that(slidelayouts, is_not(None))
    assert_that(len(slidelayouts), is_(11))


@then('I see the pptx file in the working directory')
def step_then_see_pptx_file_in_working_dir(context):
    assert_that(os.path.isfile(saved_pptx_path))
    minimum = 30000
    actual = os.path.getsize(saved_pptx_path)
    assert_that(actual, is_(greater_than(minimum)))